#!/usr/bin/env python3
"""
convert.py — turn the MGE Project & Invoice Tracker workbook into data.json
for the MGE Project & Billing Command Centre (index.html).

Usage
    python convert.py "MGE-PROJECT___INVOICE_TRACKER__SUP-UPP_.xlsx"
    python convert.py tracker.xlsx --asof 2026-10-06 --out data.json
    python convert.py tracker.xlsx --projects-sheet ASHRAF --invoices-sheet KHAN

Requires: Python 3.9+, openpyxl  (pip install openpyxl)

What it does
  * Reads the project register sheet (default "ASHRAF") and the invoice-plan
    sheet (default "KHAN"). Columns are found by their header text, so column
    order can change without breaking the script.
  * Detects the weekly status columns ("WEEK - 40", "WEEK - 41", ...) and keeps
    the latest two.
  * Detects every "<MONTH> <YEAR> PLANNED / ACTUAL" pair in the invoice sheet,
    so adding January 2027 columns just works.
  * Cleans common data-entry problems (PO values formatted as dates, 1500%
    progress, stray text in date cells) and records every fix and every
    suspicious row in a data-quality list that the dashboard displays.
"""
from __future__ import annotations

import argparse
import warnings
import datetime as dt
import json
import math
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required: pip install openpyxl")

warnings.filterwarnings("ignore", module="openpyxl")

EXCEL_EPOCH = dt.datetime(1899, 12, 30)
MONTHS = {m: i for i, m in enumerate(
    ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
     "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"], start=1)}
MON_SHORT = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# Header text (normalised) -> field name. First match wins.
PROJECT_COLS = {
    "yr": ["YEAR"], "grp": ["GROUP"], "co": ["COMPANY"], "cl": ["END-CLIENT", "END CLIENT", "CLIENT"],
    "jc": ["JOB CARD STATUS"], "so": ["ZOHO SO #", "ZOHO SO"], "ds": ["DESCRIPTION"],
    "st": ["PROJECT STATUS"], "pct": ["PROJECT %"], "lpo": ["LPO DATE"],
    "cpo": ["CUSTOMER ORDER #", "CUSTOMER ORDER"], "po": ["PO VALUE"],
    "ps": ["PS DATE"], "pf": ["PF DATE"], "as": ["ACTUAL START"], "af": ["ACTUAL FINISH"],
    "pe": ["PROJECT ENGINEER"], "se": ["SERVICE ENGINEER"],
    "rep": ["FINAL REPORT / DOCUMENTATION", "FINAL REPORT"], "wr": ["WCC REFERENCE"],
    "ws": ["WCC SUBMITTED DATE"], "wa": ["WCC APPROVED DATE"],
    "wc_ref": ["WARRANTY CERTIFICATE REFERENCE"], "sed": ["SERVICE ENTRY CREATED DATE"],
    "sen": ["SERVICE ENTRY NUMBER"], "inv": ["INVOICE NUMBER"], "ia": ["INVOICE APPROVAL"],
}
INVOICE_COLS = {
    "so": ["ZOHO SO #", "ZOHO SO"], "ds": ["PROJECT DESCRIPTION", "DESCRIPTION"], "grp": ["GROUP"],
    "cl": ["CUSTOMER ORDER #", "CUSTOMER"], "pon": ["PO NUMBER"], "po": ["PO VALUE"],
    "im": ["INVOICE MONTH"], "pm": ["PROJECTED MONTH"], "amt": ["INVOICE AMOUNT"],
    "raised": ["TOTAL INVOICE RAISED"], "bal": ["BALANCE TO INVOICE"],
    "pay": ["PAYMENT STATUS"], "age": ["PAST DUE AGE"],
}
VALID_STATUS = {"OPEN PO", "IN PROGRESS", "WORK COMPLETED", "INVOICED", "CANCELLED"}


def norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").replace("\n", " ")).strip().upper()


def blank(v) -> bool:
    return v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and not v.strip())


def txt(v) -> str:
    if blank(v):
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).replace("\xa0", "").strip()


def to_date(v):
    """Return 'YYYY-MM-DD' or None. Accepts datetimes and dd-mm-yyyy / dd/mm/yyyy text."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})\s*$", v)
        if m:
            try:
                return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
            except ValueError:
                return None
    return None


def to_num(v):
    if isinstance(v, bool) or blank(v):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    try:
        return round(float(str(v).replace(",", "").strip()), 2)
    except ValueError:
        return None


def map_columns(header_row, spec):
    idx = {}
    normed = [norm(h) for h in header_row]
    for field, names in spec.items():
        for name in names:
            if name in normed:
                idx[field] = normed.index(name)
                break
    return idx, normed


def find_header_row(ws, must_contain: str, scan: int = 15):
    for r_i, row in enumerate(ws.iter_rows(min_row=1, max_row=scan, values_only=True), start=1):
        if any(norm(c) == must_contain for c in row):
            return r_i, list(row)
    raise SystemExit(f'Sheet "{ws.title}": could not find a header row containing "{must_contain}".')


def read_projects(ws, dq):
    hr, header = find_header_row(ws, "PROJECT STATUS")
    idx, normed = map_columns(header, PROJECT_COLS)
    missing = [f for f in ("ds", "st", "po") if f not in idx]
    if missing:
        raise SystemExit(f'Sheet "{ws.title}": missing required columns {missing}.')
    week_cols = sorted([(int(m[1]), i) for i, h in enumerate(normed) if (m := re.match(r"^WEEK\s*-?\s*(\d+)$", h))])
    wk_cur = week_cols[-1] if week_cols else None
    wk_prev = week_cols[-2] if len(week_cols) > 1 else None

    get = lambda row, f: row[idx[f]] if f in idx and idx[f] < len(row) else None
    out = []
    for r_i, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
        if all(blank(c) for c in row):
            continue
        ds = txt(get(row, "ds"))
        if not ds:
            continue
        # PO value
        raw_po = get(row, "po")
        po, flag = None, ""
        if isinstance(raw_po, (dt.datetime, dt.date)):
            po = (dt.datetime.combine(raw_po, dt.time()) if isinstance(raw_po, dt.date) and not isinstance(raw_po, dt.datetime) else raw_po)
            po = (po - EXCEL_EPOCH).days
            dq.append(["PO value formatted as date", f"Row {r_i}: {ds[:45]} — recovered AED {po:,}"])
        else:
            po = to_num(raw_po)
            if po is None:
                flag = txt(raw_po) or "blank"
                dq.append(["PO value missing", f"Row {r_i}: {ds[:50]}" + (f" ({flag})" if flag != "blank" else "")])
        # progress
        pct = to_num(get(row, "pct")) or 0.0
        if pct > 1:
            dq.append(["Project %", f"Row {r_i}: {ds[:50]} — {pct*100:.0f}% read as {pct:.0f}%"])
            pct = pct / 100 if pct <= 100 else 1.0
        st = txt(get(row, "st")).upper()
        if st and st not in VALID_STATUS:
            dq.append(["Unknown project status", f"Row {r_i}: '{st}' — {ds[:45]}"])
        yr_raw = to_num(get(row, "yr"))
        rec = {
            "id": r_i, "yr": int(yr_raw) if yr_raw else None, "grp": txt(get(row, "grp")).upper(),
            "co": txt(get(row, "co")), "cl": txt(get(row, "cl")), "jc": txt(get(row, "jc")).upper(),
            "so": txt(get(row, "so")), "ds": ds, "st": st, "pct": round(pct, 3),
            "lpo": to_date(get(row, "lpo")), "cpo": txt(get(row, "cpo")), "po": po, "pof": flag,
            "ps": to_date(get(row, "ps")), "pf": to_date(get(row, "pf")),
            "as": to_date(get(row, "as")), "af": to_date(get(row, "af")),
            "pe": txt(get(row, "pe")).upper(), "se": txt(get(row, "se")).upper(),
            "rep": txt(get(row, "rep")), "wr": txt(get(row, "wr")), "ws": txt(get(row, "ws")),
            "wa": txt(get(row, "wa")), "wcr": txt(get(row, "wc_ref")), "sed": txt(get(row, "sed")),
            "sen": txt(get(row, "sen")), "inv": txt(get(row, "inv")), "ia": txt(get(row, "ia")),
            "wc": txt(row[wk_cur[1]]) if wk_cur and wk_cur[1] < len(row) else "",
            "wp": txt(row[wk_prev[1]]) if wk_prev and wk_prev[1] < len(row) else "",
        }
        if rec["yr"] is None and rec["lpo"]:
            rec["yr"] = int(rec["lpo"][:4])
        if rec["yr"] is None:
            dq.append(["Year missing", f"Row {r_i}: {ds[:50]}"])
            rec["yr"] = 0
        # stray text in date columns
        for f, label in (("pf", "PF DATE"), ("af", "ACTUAL FINISH"), ("ps", "PS DATE")):
            v = get(row, f)
            if isinstance(v, str) and not to_date(v) and v.strip().upper() not in ("TBD", "IN PROGRESS", "-", "PENDING", "N/A"):
                dq.append(["Stray text in a date column", f"Row {r_i}: '{v.strip()}' in {label} — {ds[:40]}"])
        if st == "INVOICED" and rec["jc"] == "OPEN":
            dq.append(["Job card still OPEN after invoicing", f"Row {r_i}: {rec['so']} {ds[:40]}"])
        if st == "INVOICED" and rec["inv"] in ("", "-"):
            dq.append(["Invoiced without invoice number", f"Row {r_i}: {rec['so'] or '(no SO)'} {ds[:40]}"])
        out.append(rec)
    weeks = {"cur": f"Week {wk_cur[0]}" if wk_cur else "Latest week",
             "prev": f"Week {wk_prev[0]}" if wk_prev else "Previous week"}
    return out, weeks


def read_invoices(ws, dq):
    hr, header = find_header_row(ws, "INVOICE AMOUNT")
    idx, normed = map_columns(header, INVOICE_COLS)
    # month columns: "SEPTEMBER 2026 PLANNED" / "SEPTEMBER 2026 ACTUAL"
    mcols = {}
    for i, h in enumerate(normed):
        m = re.match(r"^([A-Z]+)\s+(\d{4})\s+(PLANNED|PLAN|ACTUAL)$", h)
        if m and m[1] in MONTHS:
            key = f"{m[2]}-{MONTHS[m[1]]:02d}"
            mcols.setdefault(key, {})["plan" if m[3].startswith("PLAN") else "act"] = i
    keys = sorted(mcols)
    months = [{"key": k, "label": f"{MON_SHORT[int(k[5:]) - 1]} {k[:4]}", "start": f"{k}-01"} for k in keys]

    get = lambda row, f: row[idx[f]] if f in idx and idx[f] < len(row) else None
    out = []
    for r_i, row in enumerate(ws.iter_rows(min_row=hr + 1, values_only=True), start=hr + 1):
        ds = txt(get(row, "ds"))
        if not ds:
            continue
        cell = lambda j: to_num(row[j]) if j is not None and j < len(row) else None
        rec = {
            "id": r_i, "so": txt(get(row, "so")).replace("\n", " / "), "ds": ds, "grp": txt(get(row, "grp")).upper(),
            "cl": txt(get(row, "cl")), "pon": txt(get(row, "pon")), "po": to_num(get(row, "po")),
            "im": to_date(get(row, "im")), "pm": to_date(get(row, "pm")),
            "plan": [cell(mcols[k].get("plan")) for k in keys],
            "act": [cell(mcols[k].get("act")) for k in keys],
            "amt": to_num(get(row, "amt")), "raised": to_num(get(row, "raised")),
            "bal": to_num(get(row, "bal")), "pay": txt(get(row, "pay")).upper(), "age": txt(get(row, "age")),
        }
        if rec["po"] is not None and rec["amt"] is not None and abs(rec["amt"] - rec["po"]) > 1:
            dq.append(["Invoice amount ≠ PO value (invoice plan)",
                       f"Row {r_i}: {rec['so']} {ds[:40]} — PO {rec['po']:,.0f} vs invoice amount {rec['amt']:,.0f}"])
        if rec["pay"] == "PAID IN FULL" and (rec["bal"] or 0) > 1:
            dq.append(["Marked paid but balance still open", f"Row {r_i}: {rec['so']} {ds[:40]}"])
        out.append(rec)
    return out, months


def main():
    ap = argparse.ArgumentParser(description="Convert the MGE tracker workbook to data.json")
    ap.add_argument("xlsx", help="Path to the tracker .xlsx")
    ap.add_argument("--out", default="data.json", help="Output file (default: data.json)")
    ap.add_argument("--asof", default=None, help="Reporting date YYYY-MM-DD (default: today)")
    ap.add_argument("--projects-sheet", default="ASHRAF")
    ap.add_argument("--invoices-sheet", default="KHAN")
    a = ap.parse_args()

    asof = a.asof or dt.date.today().isoformat()
    try:
        dt.date.fromisoformat(asof)
    except ValueError:
        sys.exit("--asof must be YYYY-MM-DD")

    wb = openpyxl.load_workbook(a.xlsx, data_only=True, read_only=True)
    for s in (a.projects_sheet, a.invoices_sheet):
        if s not in wb.sheetnames:
            sys.exit(f'Sheet "{s}" not found. Sheets in the file: {wb.sheetnames}')

    dq: list[list[str]] = []
    projects, weeks = read_projects(wb[a.projects_sheet], dq)
    invoices, months = read_invoices(wb[a.invoices_sheet], dq)

    sos = {p["so"] for p in projects if p["so"]}
    pos = {p["cpo"] for p in projects if p["cpo"]}
    for v in invoices:
        if v["so"] not in sos and v["pon"] not in pos:
            dq.append(["Invoice-plan line not found in the register", f"{v['so'] or '(no SO)'} {v['ds'][:50]}"])

    data = {
        "asof": asof,
        "generated": dt.datetime.now().isoformat(timespec="seconds"),
        "source": Path(a.xlsx).name,
        "sheets": {"projects": a.projects_sheet, "invoices": a.invoices_sheet},
        "weeks": weeks,
        "months": months,
        "projects": projects,
        "invoices": invoices,
        "dq": dq,
    }
    Path(a.out).write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"Wrote {a.out}: {len(projects)} job lines, {len(invoices)} invoice-plan lines, "
          f"{len(months)} plan months ({months[0]['label'] if months else '-'} to {months[-1]['label'] if months else '-'}), "
          f"{len(dq)} data-quality notes, as of {asof}.")


if __name__ == "__main__":
    main()
